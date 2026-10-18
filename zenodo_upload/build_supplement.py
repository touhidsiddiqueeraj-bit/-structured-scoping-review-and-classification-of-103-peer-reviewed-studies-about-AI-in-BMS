#!/usr/bin/env python3
"""Build Supplementary_Material.docx — robustness material consolidated out of
the main text (interpretation rules, label-noise scenarios, MDE analysis,
recall-audit construction, abstract-level probes).

Run from repo root:  python3 build_supplement.py
Then convert:        soffice --headless --convert-to pdf Supplementary_Material.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
import paper_part1  # for the title

OUT = os.path.join(HERE, 'Supplementary_Material.docx')

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10.5)

def para(text, size=10.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_after=8, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    # minimal *italic* markup
    parts = text.split('*')
    for i, chunk in enumerate(parts):
        if not chunk:
            continue
        r = p.add_run(chunk)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic or (i % 2 == 1)
    return p

def h1(text):
    para(text, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)

def h2(text):
    para(text, size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

def artifacts(text):
    para(text, size=9.5, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)

# ------------------------------------------------------------------ title
para('Supplementary Material', size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para(paper_part1.TITLE, size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

para('This supplement consolidates the robustness material that supports the main text: the '
     'interpretation rules that govern every quantitative claim (S1), the label-noise scenario '
     'envelope for the validation counts (S2), the minimum-detectable-effect analysis of the '
     'method-family contrast (S3), the construction and limits of the recall audit (S4), and the '
     'two abstract-level probes (S5). Every section states the released artifacts it derives from; '
     'the complete evidence base and analysis pipeline are deposited at Zenodo, DOI '
     '10.5281/zenodo.23436202, with the GitHub repository of the same name as a development '
     'mirror. The main text carries one-sentence summaries of each item with pointers here, so '
     'the supplement can be read alongside Sections 2.2, 2.3, and 3.4 of the paper.')

# ------------------------------------------------------------------ S1
h1('S1. Interpretation rules for every quantitative claim in the paper')
para('Five rules govern how the paper\u2019s numbers may be read. They are stated here in full '
     'because the main text invokes them by number.')
para('*Rule 1 (subset-scoped claims).* Anything coded from full texts \u2014 function, method '
     'family, chemistry, data source, validation protocol, baseline screen \u2014 describes the '
     '103-study classified subset only and is never generalized to the field. Headlines that read '
     '\u2018the subset\u2019s\u2019 or \u2018within the classified subset\u2019 invoke this rule.')
para('*Rule 2 (pool-scoped claims).* Tier-1 counts (799 records) measure a capped 30-query '
     'Crossref retrieval, not a census. Per-year counts are observed counts inside that cap; '
     'query-theme tag shares are read for ordering, not exact rates; the 2024 window is partially '
     'covered (no query pass was dedicated to January\u2013May 2024); and the recall audit of S4 '
     '(16\u201317% capture of the reviews\u2019 own cited literature) bounds external validity.')
para('*Rule 3 (earliest means earliest in the classified subset).* Recall is bounded, so '
     '\u2018first\u2019 and \u2018earliest\u2019 claims are subset-first by construction and are '
     'phrased accordingly.')
para('*Rule 4 (non-significance is not equivalence).* The physics-informed versus data-driven '
     'contrast is underpowered (minimum detectable effect roughly 30\u201335 percentage points, '
     'S3); \u2018no detectable difference\u2019 is reported as exactly that.')
para('*Rule 5 (curation conditions the subset).* Tier-2 selection is quota- and '
     'citation-ranking-driven (main text, Section 2.3), so distributions inside the 103 partly '
     'reflect those choices and are reported as corpus composition.')
artifacts('Invoked at: main text Sections 2.2, 2.3, 2.6, 3.2\u20133.4, and the Figure 5 caption. '
          'Released artifacts: selection.json, screening_log.csv, pool.csv, recall_audit.json.')

# ------------------------------------------------------------------ S2
h1('S2. Label-noise scenario envelope for the validation counts')
para('The strong/weak validation counts of the main text rest on a single full-text coding by the '
     'lead author (the reliability study of Section 2.5 measures abstract-level reproducibility, '
     'not full-text agreement). To quantify how much those counts depend on the labels, a '
     'label-noise scenario analysis perturbs them directly. The analysis is a scenario envelope, '
     'not an error model.')
para('*Design.* In each of 4,000 seeded Monte Carlo draws (seed 7), a label-flip rate is drawn '
     'from a Normal distribution centered on an observed coder-versus-label error rate with '
     'standard deviation 0.02, truncated at zero; the drawn rate then flips strong/non-strong '
     'bucket labels across the 88 primary papers, and both the overall strong share and the '
     'physics-versus-other Fisher exact test are recomputed.')
para('*Scenario 1 (symmetric).* The rate is centered on the coders\u2019 overall bucket error '
     'rate against the released labels (0.37 and 0.27 for the two coders; mean 0.32), and flips '
     'labels in either direction. Symmetric random flipping mechanically attenuates contrasts '
     'toward the null.')
para('*Scenario 2 (strong-to-weak, worst-case bound).* The actually observed coder-label '
     'disagreements are directional and concentrated on the strong labels: 16 of the 19 bucket '
     'disagreements have the coder assigning weak where the released label says strong (the '
     'declared conservative fallback of Section 2.5), 13 of the 30 sampled papers are '
     'strong-labelled, and each coder recoded 8 of those 13 weak \u2014 a strong-label-specific '
     'rate of 0.62 (16 of 26 pooled coder-labels). Scenario 2 therefore flips strong labels to '
     'weak only, at a rate drawn from a Normal centered on that strong-label-specific rate. '
     'Because the coders worked from abstracts under a declared weak-if-unreported rule, that '
     'rate is an upper bound on plausible label error, and Scenario 2 is reported as a worst-case '
     'bound, not an estimate.')
para('*Results.* Under the symmetric scenario the strong-protocol count spans 35\u201352 of the '
     '88 primary papers (39.8\u201359.1%) and the family contrast remains non-significant in '
     '94.4% of draws. Under the strong-to-weak worst-case bound the count falls to about 16 of 88 '
     '(\u224818%) at the observed recode rate, with a scenario envelope of 10\u201323 '
     '(11.4\u201326.1%), and the contrast remains non-significant in 97.6% of draws. Read '
     'together, the two scenarios bracket the single-coder counts around the documented '
     'disagreement structure \u2014 under either, the physics-versus-other contrast shows no '
     'detectable difference and remains underpowered in the sense quantified in S3 \u2014 and the '
     'worst-case bound says only how low the strong-protocol count could sit if abstract-level '
     'recoding were taken at face value, not where it sits.')
artifacts('Released artifacts: sensitivity_label_noise.py and sensitivity_label_noise.json; '
          'gold_labels_30.csv (the 30-sample labels); coder_B_30.csv, coder_C_30.csv, '
          'agreement_report.csv, dispute_items.csv, and the hash-tracked reliability log under '
          'corpus/reliability/.')

# ------------------------------------------------------------------ S3
h1('S3. Minimum-detectable-effect analysis of the method-family contrast')
para('The main-text contrast \u2014 8 of 22 physics-informed/hybrid primary papers with a strong '
     'protocol versus 34 of 66 others (protocol-only definition) \u2014 is accompanied by a '
     'power analysis so that non-significance is not misread as equivalence.')
para('*Method.* Analytic two-proportion power, validated against Fisher-exact Monte Carlo '
     'simulation, is computed for the observed arm sizes and a 51.5% baseline strong-protocol '
     'rate. Under the broader definition of stringency (which also counts '
     'experimental-hardware demonstrations) the same analysis is repeated for the 13-of-22 versus '
     '49-of-66 contrast.')
para('*Results.* With 22 versus 66 papers, 80% power at \u03b1 = 0.05 requires a difference of '
     'roughly 30\u201335 percentage points between the families (odds ratio \u22640.26 or '
     '\u22654.0); the observed 15.2-point gap is detectable with only about 20\u201325% power. '
     'The broader-definition contrast is similar. The two arms\u2019 Wilson 95% intervals overlap '
     'almost completely (19.7\u201357.0% versus 39.7\u201363.2%). The honest statement is that '
     'no difference is detectable at this scale, and that meaningful differences in either '
     'direction \u2014 including a physics prior halving or doubling the odds of a strong '
     'protocol \u2014 remain compatible with these counts.')
artifacts('Released artifacts: power_analysis.py and power_analysis.json.')

# ------------------------------------------------------------------ S4
h1('S4. Recall-audit construction, yields, and limits')
para('The main text reports that the 1,550-record retrieval contains 16\u201317% of the '
     'in-scope literature cited by this review\u2019s own included reviews. This section states '
     'how that number is built.')
para('*Denominator.* The deposited reference lists of the 15 included reviews and the six seed '
     'studies \u2014 2,508 references in total \u2014 are restricted to references dated 2020 or '
     'later (1,367 are dated before 2020 and fall outside the retrieval window by design), then '
     'to references that pass the same title gate as the screening pipeline and carry matchable '
     'DOIs: 649 counted with multiplicity, 586 unique.')
para('*Numerator and result.* The 1,550-record retrieval contains 112 of the 649 (17.3% '
     'per-source) and 96 of the 586 unique references (16.4%); all 96 are Tier-1 eligible. '
     'Per-query yields are design-capped, not saturated: every pass returned its top 20 records '
     'only, so the 30 query families are capped at 100 records each; the five passes yielded '
     '3,000 fetched rows, 2,973 after document-type, year, and title-completeness filters, and '
     '1,550 unique records after deduplication.')
para('*Limits.* The denominator is a citation-conditioned sample of what the reviews happened to '
     'cite; the in-scope gate runs on titles; and references without deposited DOIs cannot be '
     'matched. The audit is therefore a floor on external validity for theme frequencies, '
     'per-family shares, and \u2018earliest\u2019 claims, which is why all such claims in the '
     'paper are subset- or pool-scoped (S1, rules 2 and 3). The recall miss documented in main-'
     'text Section 8.4 (the TEC 2025 GPT-4 estimator) fits this pattern exactly: it is cited by '
     'none of the included reviews either, so no reference-list audit would have surfaced it.')
artifacts('Released artifacts: recall_audit.py, recall_audit.json, and the cached Crossref '
          'responses under raw_cache/.')

# ------------------------------------------------------------------ S5
h1('S5. Abstract-level probes')
h2('S5.1 Representativeness probe: can the pool carry the validation audit?')
para('*Question.* Does the curated subset overstate validation evidence relative to the pool it '
     'was drawn from \u2014 at the one level where the 799-record pool is observable, its '
     'indexed abstracts?')
para('*Design.* A seed-pinned uniform sample of 40 eligible records was drawn (released with its '
     'codes); 23 abstracts were retrievable from public indexes (Elsevier-withheld abstracts '
     'dominate the 17 missing). Each was coded under a three-value rubric: isolation or '
     'deployment protocol stated; evaluation stated without isolation; nothing stated. The pool '
     'sample\u2019s record mix: 16 primary-ML, 2 model-based, 3 review, 2 other.')
para('*Result.* None of the 18 primary-looking pool records states an isolation or deployment '
     'protocol in its abstract (0 of 18), against 5 of the 38 abstract-screenable subset papers '
     '(13%); the difference is not significant (Fisher\u2019s exact p = 0.16) and both rates are '
     'near zero. The finding is negative in a useful way: indexed abstracts almost never carry '
     'enough protocol information to code validation strength at pool scale, in either tier, so '
     'the 48% strong-protocol share cannot be promoted from a curated-subset statistic to a field '
     'estimate without full-text access \u2014 and the scarcity of protocol statements in '
     'abstracts is itself part of what the reporting checklist of main-text Section 9.9 responds '
     'to.')
h2('S5.2 Baseline-comparison screen: do ML studies benchmark against model-based methods?')
para('*Question.* The review had audited validation protocols but not whether ML studies compare '
     'against model-based or classical reference methods at all \u2014 a Kalman-family filter '
     '(EKF/UKF), an observer, an equivalent-circuit or physics-based model, coulomb counting, or '
     'a classical regressor such as Gaussian-process regression.')
para('*Design.* A seventh, lighter classification field (added when the review\u2019s own '
     'checklist turned on its authors) records, from indexed abstracts, whether such a reference '
     'method is named as a comparator: values none, mentioned, and stated same-data comparison. '
     'It is coded from abstracts, not full texts; the released rubric (codebook.md) also defines '
     'the full-text version of the field for future work.')
para('*Result.* Of the 38 abstract-screenable primary papers, 32 (84%) name no model-based or '
     'classical baseline, 2 name one without stating the comparison, and 4 (11%) state a '
     'same-data quantitative comparison \u2014 against an SVR, a 0.5C CCCV protocol, an '
     'equivalent-circuit/electrochemical model, and a rule-based controller, respectively. '
     'These are lower bounds on comparison practice (full texts may compare what abstracts '
     'omit) but exact measures of comparison visibility, and visibility is what a reader '
     'scanning the literature actually has: at abstract level, the field\u2019s answer to '
     '\u2018does the learned estimator beat the model-based baseline on the same data?\u2019 '
     'is, four times in five, unstated.')
artifacts('Released artifacts: abstract_probe.py, abstract_probe_raw.json, '
          'abstract_probe_codes.py, abstract_probe.json; rubric in codebook.md.')

# ------------------------------------------------------------------ close
para('All scripts run from the repository root against the released data and reproduce their '
     'outputs deterministically (seeded where stochastic). Deposit of record: Zenodo, DOI '
     '10.5281/zenodo.23436202; development mirror: the GitHub repository of the same name.',
     size=9.5, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)

doc.save(OUT)
print('saved', OUT)
