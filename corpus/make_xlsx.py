#!/usr/bin/env python3
"""Build the corpus spreadsheet: 103 verified papers with clickable DOI links."""
import os
import json, csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---- local replacements for the previous skill-template helpers ----
_TITLE_FONT = Font(name='Calibri', size=14, bold=True, color='1F4E79')
_HDR_FILL = PatternFill('solid', fgColor='1F4E79')
_HDR_FONT = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
_THIN = Side(style='thin', color='B0B0B0')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

def setup_sheet(ws, title, last_col, start_row=2):
    ws.cell(row=start_row, column=2, value=title).font = _TITLE_FONT
    ws.sheet_view.showGridLines = False

def style_header_row(ws, row_num, col_start, col_end):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row_num, column=c)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def style_data_row(ws, row_num, col_start, col_end, row_index=0):
    band = 'F2F6FA' if row_index % 2 else 'FFFFFF'
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row_num, column=c)
        cell.border = _BORDER
        if not cell.font or not cell.font.bold:
            cell.font = Font(name='Calibri', size=9)
        cell.fill = PatternFill('solid', fgColor=band)
        cell.alignment = Alignment(vertical='top', wrap_text=True)

def auto_fit_columns(ws, min_width=8, max_width=44, header_row=4, data_start_row=5):
    for col in range(2, ws.max_column + 1):
        letter = get_column_letter(col)
        width = min_width
        for r in range(header_row, ws.max_row + 1):
            v = ws.cell(row=r, column=col).value
            if v is None:
                continue
            w = min(max_width, max(min_width, int(1.05 * min(len(str(v)), 60)) + 2))
            width = max(width, w)
        ws.column_dimensions[letter].width = width

def auto_fit_row_heights(ws, header_row=4, data_start_row=5):
    for r in range(data_start_row, ws.max_row + 1):
        ws.row_dimensions[r].height = 14


CORPUS = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'corpus/papers_final.json')))

wb = Workbook()

# ================= Sheet 1: Papers =================
ws = wb.active
ws.title = "Papers"

HEADERS = ['#', 'Title', 'Authors', 'Year', 'Journal', 'Link (DOI)', 'DOI',
           'BMS Function', 'Method Family', 'Chemistry', 'Data Source',
           'Validation', 'Citations', 'Key Finding', 'Key Limitation']
last_col = len(HEADERS) + 1
setup_sheet(ws, title="Machine Learning in Battery Management Systems \u2014 Verified Corpus (103 papers, 2019\u20132026)", last_col=last_col)

for c, h in enumerate(HEADERS, start=2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

# sort by function, then year desc, then citations desc
ORDER = ['State estimation (SOC)', 'Health estimation (SOH)', 'Prognostics (RUL)',
         'Charging control', 'Thermal management', 'Fault & safety',
         'Cell balancing & ECM', 'Cross-cutting']
rows = sorted(CORPUS, key=lambda r: (ORDER.index(r['area']), -r['year'], -r['cites']))

DATA_START = 5
for i, r in enumerate(rows):
    rn = DATA_START + i
    url = 'https://doi.org/' + r['doi']
    vals = [i + 1,
            r['title_cr'].replace('&amp;', '&'),
            (r['authors'] or '').replace(';', ',').replace(', et al.', ' et al.'),
            r['year'], r['journal'].replace('&amp;', '&'),
            url, r['doi'], r['area'], r['method'], r['chem'], r['data'],
            r['valid'], r['cites'],
            r['finding'][0].upper() + r['finding'][1:], r['limit'][0].upper() + r['limit'][1:]]
    for c, v in enumerate(vals, start=2):
        ws.cell(row=rn, column=c, value=v)
    style_data_row(ws, row_num=rn, col_start=2, col_end=last_col, row_index=i)
    # hyperlink on the Link column (col index of 'Link (DOI)' = 7)
    link_cell = ws.cell(row=rn, column=7)
    link_cell.hyperlink = url
    link_cell.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
    # numeric right-align
    ws.cell(row=rn, column=2).alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(row=rn, column=5).alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(row=rn, column=14).alignment = Alignment(horizontal='right', vertical='center')

ws.auto_filter.ref = f"B4:{get_column_letter(last_col)}{DATA_START + len(rows) - 1}"
ws.freeze_panes = "D5"  # keep header + #/title visible when scrolling
auto_fit_columns(ws, min_width=8, max_width=44, header_row=4, data_start_row=DATA_START)
# widen long-text columns explicitly
for col, w in {'C': 46, 'D': 30, 'F': 26, 'G': 30, 'I': 22, 'J': 20, 'O': 42, 'P': 40}.items():
    ws.column_dimensions[col].width = w
auto_fit_row_heights(ws, header_row=4, data_start_row=DATA_START)
cap = ws.cell(row=DATA_START + len(rows) + 1, column=2,
              value='All 103 bibliographic records verified against the Crossref registry (api.crossref.org) by direct DOI resolution on 5 September 2026. Citations = Crossref is-referenced-by-count at retrieval. Classification fields per the paper\u2019s Section 2.5.')
cap.font = Font(name='Calibri', size=9, italic=True, color='595959')
ws.merge_cells(start_row=DATA_START + len(rows) + 1, start_column=2,
               end_row=DATA_START + len(rows) + 1, end_column=10)

# ================= Sheet 2: Summary =================
ws2 = wb.create_sheet("Summary")
from collections import Counter, defaultdict
by_area = Counter(r['area'] for r in CORPUS)
by_year = Counter(r['year'] for r in CORPUS)
by_method = Counter(G if False else r['method'] for r in CORPUS)
by_data = Counter(r['data'] for r in CORPUS)

setup_sheet(ws2, title="Corpus Summary \u2014 103 peer-reviewed studies", last_col=4)

def block(ws2, start_row, title, counter, order=None, label='Category'):
    ws2.cell(row=start_row, column=2, value=title).font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
    ws2.cell(row=start_row + 1, column=2, value=label)
    ws2.cell(row=start_row + 1, column=3, value='Papers')
    style_header_row(ws2, row_num=start_row + 1, col_start=2, col_end=3)
    items = order if order else sorted(counter.items(), key=lambda kv: -kv[1])
    if order:
        items = [(k, counter.get(k, 0)) for k in order]
    for i, (k, v) in enumerate(items):
        rn = start_row + 2 + i
        ws2.cell(row=rn, column=2, value=k)
        c = ws2.cell(row=rn, column=3, value=v)
        style_data_row(ws2, row_num=rn, col_start=2, col_end=3, row_index=i)
        c.alignment = Alignment(horizontal='right', vertical='center')
    return start_row + 2 + len(items) + 1

nxt = block(ws2, 4, 'By BMS management function', by_area, order=ORDER)
nxt = block(ws2, nxt, 'By publication year', by_year, order=sorted(by_year))
nxt = block(ws2, nxt, 'By method family', by_method)
nxt = block(ws2, nxt, 'By primary data source', by_data)
auto_fit_columns(ws2, min_width=8, max_width=44, header_row=5, data_start_row=6)
auto_fit_row_heights(ws2, header_row=5, data_start_row=6)

# ================= Sheet 3: Conference sensitivity =================
ws3 = wb.create_sheet("Conference sensitivity")
C3 = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'corpus/conference_verified.json')))
SENS = csv.DictReader(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'corpus/conference_sensitivity.csv')))
val = {r['doi']: r['validation_evidence'] for r in SENS}
fn = {r['doi']: r['function'] for r in SENS}
H3 = ['#', 'Title', 'Authors', 'Year', 'Venue', 'Link (DOI)', 'DOI', 'Function',
      'Method (short)', 'Validation evidence stated in indexed abstract']
last3 = len(H3) + 1
setup_sheet(ws3, title="Conference Sensitivity Set \u2014 14 proceedings papers (CrossRef-verified; outside the 103-paper corpus)", last_col=last3)
for c, h in enumerate(H3, start=2):
    ws3.cell(row=4, column=c, value=h)
style_header_row(ws3, row_num=4, col_start=2, col_end=last3)
for i, r in enumerate(sorted(C3, key=lambda x: (fn.get(x['doi'], ''), x.get('year') or 0))):
    rn = 5 + i
    url = 'https://doi.org/' + r['doi']
    vals = [i + 1, r['title_cr'].replace('&amp;', '&'),
            (r['authors'] or '').replace(';', ',').replace(', et al.', ' et al.'),
            r['year'], r['journal'].replace('&amp;', '&'), url, r['doi'],
            fn.get(r['doi'], ''), r['method'], val.get(r['doi'], '')]
    for c, v in enumerate(vals, start=2):
        ws3.cell(row=rn, column=c, value=v)
    style_data_row(ws3, row_num=rn, col_start=2, col_end=last3, row_index=i)
    lc = ws3.cell(row=rn, column=7)
    lc.hyperlink = url
    lc.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
    ws3.cell(row=rn, column=2).alignment = Alignment(horizontal='right', vertical='center')
    ws3.cell(row=rn, column=5).alignment = Alignment(horizontal='right', vertical='center')
ws3.auto_filter.ref = f"B4:{get_column_letter(last3)}{5 + len(C3) - 1}"
ws3.freeze_panes = "D5"
auto_fit_columns(ws3, min_width=8, max_width=44, header_row=4, data_start_row=5)
for col, w in {'C': 46, 'D': 30, 'F': 26, 'G': 30, 'J': 22, 'K': 40}.items():
    ws3.column_dimensions[col].width = w
auto_fit_row_heights(ws3, header_row=4, data_start_row=5)
cap3 = ws3.cell(row=5 + len(C3) + 1, column=2,
                value='Curated for the conference-path sensitivity analysis (paper Section 3.5): venue-scoped Crossref queries over ACC/ECC/CCDC/APEC/VPPC/ITEC, proceedings articles 2019\u20132025, ranked by citations, verified by direct DOI resolution on 5 September 2026. Kept outside the 103-paper corpus by design.')
cap3.font = Font(name='Calibri', size=9, italic=True, color='595959')
ws3.merge_cells(start_row=5 + len(C3) + 1, start_column=2, end_row=5 + len(C3) + 1, end_column=9)

wb.properties.creator = "Z.ai"
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'ML-BMS-Corpus.xlsx')
wb.save(out)
print('saved', out, '| papers:', len(rows))
